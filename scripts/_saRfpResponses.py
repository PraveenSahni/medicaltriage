import openpyxl
import re

P = r"C:\Users\PraveenSAHNI\OneDrive - IRISSTAR TECHNOLOGIES\000.PSDownload\NFR_COTS_CSQ_v8.3-edc758c5-2b70-496a-9db4-947af34a67a4.xlsx"

PREFIX = {
    "Yes": "As IST's Solution Architect for this bid, I confirm the proposed IST Health Tele-Triage platform fully meets this requirement.",
    "Partial": "As IST's Solution Architect for this bid, I confirm the proposed solution partially meets this requirement today, with a clear, disclosed path to full compliance.",
    "No": "As IST's Solution Architect for this bid, I confirm the proposed solution does not currently meet this requirement, and I have set out the honest reason and remediation path below.",
    "N/A": "As IST's Solution Architect for this bid, I confirm this requirement is not applicable to the proposed architecture, for the reason set out below.",
    "Other": "As IST's Solution Architect for this bid, this item remains open pending the input noted below and cannot yet be given a definitive Yes/No/Partial response.",
}

FALLBACK_NO_REMARK = {
    "Yes": "The proposed solution meets this requirement as implemented in the current architecture.",
    "Partial": "The proposed solution partially meets this requirement; a fuller assessment and remediation plan will follow as part of contract mobilization.",
    "No": "The proposed solution does not currently meet this requirement; IST will address this as part of the implementation plan if awarded.",
    "N/A": "This requirement falls outside the scope of the proposed architecture.",
    "Other": "This item requires further internal review before IST can provide a definitive response.",
}

ALREADY_TAGGED = re.compile(r"^As IST's Solution Architect")


def build_response(compliance, remark):
    compliance = (compliance or "").strip()
    prefix = PREFIX.get(compliance)
    if prefix is None:
        # Unrecognized/non-standard compliance string - treat conservatively, don't guess.
        prefix = "As IST's Solution Architect for this bid, our position on this requirement is set out below."
    remark = (remark or "").strip()
    if ALREADY_TAGGED.match(remark):
        return remark  # already converted in a prior run - leave untouched (idempotent)
    if not remark:
        remark = FALLBACK_NO_REMARK.get(compliance, "")
    return f"{prefix} {remark}".strip()


def process_ai_or_nfr_or_ux(ws, id_col=1, compliance_col=6, remark_col=7):
    updated = 0
    for row in ws.iter_rows(min_row=3):
        row_id = row[id_col - 1].value
        if not row_id or not isinstance(row_id, str):
            continue
        compliance_cell = row[compliance_col - 1]
        remark_cell = row[remark_col - 1]
        new_val = build_response(compliance_cell.value, remark_cell.value)
        remark_cell.value = new_val
        updated += 1
    return updated


def process_cloud_csq(ws, id_col=4, compliance_col=9, remark_col=10):
    updated = 0
    for row in ws.iter_rows(min_row=11):
        row_id = row[id_col - 1].value
        if not row_id or not isinstance(row_id, str) or "." not in row_id:
            continue
        compliance_cell = row[compliance_col - 1]
        remark_cell = row[remark_col - 1]
        new_val = build_response(compliance_cell.value, remark_cell.value)
        remark_cell.value = new_val
        updated += 1
    return updated


wb = openpyxl.load_workbook(P, data_only=False)

counts = {}
counts["Non Functional Req"] = process_ai_or_nfr_or_ux(wb["Non Functional Req"])
counts["UX"] = process_ai_or_nfr_or_ux(wb["UX"])
counts[" AI"] = process_ai_or_nfr_or_ux(wb[" AI"])
counts["Cloud CSQ"] = process_cloud_csq(wb["Cloud CSQ"])

wb.save(P)

with open(r"C:\AiMlTriage\sa_response_run_log.txt", "w", encoding="utf-8") as f:
    total = 0
    for k, v in counts.items():
        f.write(f"{k}: {v} rows updated\n")
        total += v
    f.write(f"TOTAL: {total} rows updated\n")

print("done")

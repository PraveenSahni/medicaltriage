import openpyxl
import re

path = r"C:\Users\PraveenSAHNI\OneDrive - IRISSTAR TECHNOLOGIES\000.PSDownload\NFR_COTS_CSQ_v8.3-edc758c5-2b70-496a-9db4-947af34a67a4.xlsx"
wb = openpyxl.load_workbook(path)

def classify_response(v):
    if v is None:
        return "blank"
    s = str(v)
    stripped = s.strip()
    if stripped != s:
        flag = "WHITESPACE"
    else:
        flag = ""
    if stripped == "Yes":
        return "yes" + flag
    if "inherited" in stripped.lower() and "yes" in stripped.lower():
        return "yes_inherited" + flag
    if stripped.upper().startswith("N/A"):
        return "na" + flag
    if stripped == "Partial":
        return "partial" + flag
    if stripped == "No":
        return "no" + flag
    if "needs" in stripped.lower():
        return "other_needs" + flag
    return "unrecognized:" + repr(s)

def analyze_nfr_shape(ws, name, mandatory_col=5, compliance_col=6, id_col=1, header_row=2, start_row=3):
    total_physical = ws.max_row
    header_rows = start_row - 1
    blank = no_id = 0
    mandatory_rows = []
    optional_rows = []
    for row in range(start_row, ws.max_row + 1):
        rid = ws.cell(row=row, column=id_col).value
        mand = ws.cell(row=row, column=mandatory_col).value
        comp = ws.cell(row=row, column=compliance_col).value
        if rid is None and mand is None and comp is None:
            blank += 1
            continue
        if rid is None:
            no_id += 1
            continue
        cls = classify_response(comp)
        entry = (row, rid, mand, comp, cls)
        if mand == "Yes":
            mandatory_rows.append(entry)
        else:
            optional_rows.append(entry)
    return dict(tab=name, total_physical=total_physical, header_rows=header_rows,
                blank=blank, no_id=no_id, mandatory=mandatory_rows, optional=optional_rows)

def analyze_csq_shape(ws, name, mandatory_col=6, compliance_col=9, id_col=4, start_row=11):
    total_physical = ws.max_row
    header_rows = start_row - 1
    blank = no_id = 0
    mandatory_rows = []
    optional_rows = []
    for row in range(start_row, ws.max_row + 1):
        rid = ws.cell(row=row, column=id_col).value
        mand = ws.cell(row=row, column=mandatory_col).value
        comp = ws.cell(row=row, column=compliance_col).value
        if rid is None and mand is None and comp is None:
            blank += 1
            continue
        if rid is None:
            no_id += 1
            continue
        cls = classify_response(comp)
        entry = (row, rid, mand, comp, cls)
        if mand == "Mandatory":
            mandatory_rows.append(entry)
        else:
            optional_rows.append(entry)
    return dict(tab=name, total_physical=total_physical, header_rows=header_rows,
                blank=blank, no_id=no_id, mandatory=mandatory_rows, optional=optional_rows)

def summarize(rows):
    counts = {}
    for (row, rid, mand, comp, cls) in rows:
        base = cls.replace("WHITESPACE", "")
        counts[base] = counts.get(base, 0) + 1
    return counts

results = {}
results["Non Functional Req"] = analyze_nfr_shape(wb["Non Functional Req"], "Non Functional Req")
results["UX"] = analyze_nfr_shape(wb["UX"], "UX")
results["AI"] = analyze_nfr_shape(wb[" AI"], "AI")
results["Cloud CSQ"] = analyze_csq_shape(wb["Cloud CSQ"], "Cloud CSQ")

grand_mandatory_yes = grand_mandatory_na = grand_mandatory_total = 0
grand_optional_yes = grand_optional_na = grand_optional_total = 0
grand_ids = set()
dupes = []

for tabname, r in results.items():
    print(f"\n=== {tabname} ===")
    print(f"total_physical_rows={r['total_physical']} header_rows={r['header_rows']} blank_rows={r['blank']} rows_without_id={r['no_id']}")
    m_counts = summarize(r["mandatory"])
    o_counts = summarize(r["optional"])
    print(f"mandatory_rows={len(r['mandatory'])} breakdown={m_counts}")
    print(f"non_mandatory_rows={len(r['optional'])} breakdown={o_counts}")

    m_yes = m_counts.get("yes", 0) + m_counts.get("yes_inherited", 0)
    m_na = m_counts.get("na", 0)
    m_total = len(r["mandatory"])
    m_denom = m_total - m_na
    print(f"MANDATORY scored denom={m_denom} yes={m_yes} pct={100*m_yes/m_denom if m_denom else 0:.1f}%")

    o_yes = o_counts.get("yes", 0) + o_counts.get("yes_inherited", 0)
    o_na = o_counts.get("na", 0)
    o_total = len(r["optional"])
    o_denom = o_total - o_na
    print(f"NON-MANDATORY scored denom={o_denom} yes={o_yes} pct={100*o_yes/o_denom if o_denom else 0:.1f}%")

    all_yes = m_yes + o_yes
    all_na = m_na + o_na
    all_total = m_total + o_total
    all_denom = all_total - all_na
    print(f"ALL (this tab) scored denom={all_denom} yes={all_yes} pct={100*all_yes/all_denom if all_denom else 0:.1f}%")

    grand_mandatory_yes += m_yes
    grand_mandatory_na += m_na
    grand_mandatory_total += m_total
    grand_optional_yes += o_yes
    grand_optional_na += o_na
    grand_optional_total += o_total

    for entry_list in (r["mandatory"], r["optional"]):
        for (row, rid, mand, comp, cls) in entry_list:
            key = (tabname, rid)
            if key in grand_ids:
                dupes.append(key)
            grand_ids.add(key)
            if cls.startswith("unrecognized"):
                print(f"  UNRECOGNIZED: row={row} id={rid} value={comp!r}")

print("\n\n=== COMBINED TOTALS ===")
gm_denom = grand_mandatory_total - grand_mandatory_na
print(f"Mandatory: total={grand_mandatory_total} na={grand_mandatory_na} denom={gm_denom} yes={grand_mandatory_yes} pct={100*grand_mandatory_yes/gm_denom:.2f}%")
go_denom = grand_optional_total - grand_optional_na
print(f"Non-mandatory: total={grand_optional_total} na={grand_optional_na} denom={go_denom} yes={grand_optional_yes} pct={100*grand_optional_yes/go_denom:.2f}%" if go_denom else "Non-mandatory: n/a")
all_total = grand_mandatory_total + grand_optional_total
all_na = grand_mandatory_na + grand_optional_na
all_yes = grand_mandatory_yes + grand_optional_yes
all_denom = all_total - all_na
print(f"OVERALL: total_scored_rows={all_total} na={all_na} denom={all_denom} yes={all_yes} pct={100*all_yes/all_denom:.2f}%")

print(f"\nDuplicate (tab,id) keys found: {dupes if dupes else 'NONE'}")
